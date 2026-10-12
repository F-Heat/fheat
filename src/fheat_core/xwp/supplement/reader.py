"""Read the supplement files and merge them into the plan F|Heat built.

Merge rules (apply_supplement):

* an object whose uuid exists in the plan overrides it: every attribute the
  planner filled in replaces F|Heat's value (lists as a whole); empty cells
  keep F|Heat's value; it counts as overridden only if a value changed;
* other objects are added; a missing ``gehoertZuPlan``/``gehoertZuBereich``
  is set to the plan and its area. An added object that still lacks a
  mandatory attribute is not taken over and reported; a placeholder of the
  template the planner left incomplete only counts as open;
* a second WP_Plan is not taken over.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from fheat_core.xwp.decode.tables import from_tables
from fheat_core.xwp.document import XwpPlanDocument
from fheat_core.xwp.model.base import REF, Issue, Ref, XwpFeature, validate
from fheat_core.xwp.supplement.template import INDEX_SHEET


@dataclass
class Supplement:
    """Objects read from the supplement files, with the findings of reading them."""

    objects: list[XwpFeature] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)


@dataclass
class MergeResult:
    overridden: int = 0
    added: int = 0
    skipped: int = 0
    issues: list[Issue] = field(default_factory=list)
    open_placeholders: int = 0
    overridden_ids: set[str] = field(default_factory=set)
    added_ids: set[str] = field(default_factory=set)


def read_supplement(gpkg: str | Path | None = None, xlsx: str | Path | None = None, model=None) -> Supplement:
    """Read the GeoPackage and/or the Excel supplement."""
    if model is None:
        from fheat_core.xwp.model import current as model
    tables: dict[str, list[dict]] = {}
    if gpkg is not None:
        from fheat_core.xwp.decode.gpkg import read_tables

        tables.update(read_tables(gpkg))
    if xlsx is not None:
        for name, rows in read_xlsx_tables(xlsx).items():
            tables.setdefault(name, [])
            tables[name] = tables[name] + rows
    tables = {name: rows for name, rows in tables.items() if not _is_reference(name, model)}
    objects, issues = from_tables(tables, model)
    return Supplement(objects, issues)


def _is_reference(name: str, model) -> bool:
    """Code lists and metadata are written for reference only."""
    from fheat_core.xwp.encode.gpkg import METADATA_TABLE, table_base

    return name == METADATA_TABLE or name in {table_base(e) for e in model.ENUMS}


def read_xlsx_tables(path: str | Path) -> dict[str, list[dict]]:
    """Rows per table name from the Excel template (sheets mapped via the index sheet)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if INDEX_SHEET not in wb.sheetnames:
            raise ValueError(f"{path}: Blatt {INDEX_SHEET} fehlt; ist das eine F|Heat-Ergänzungsvorlage?")
        mapping = {}
        for row in wb[INDEX_SHEET].iter_rows(min_row=2, values_only=True):
            if row and row[0] and row[1]:
                mapping[str(row[0])] = str(row[1])
        tables: dict[str, list[dict]] = {}
        for sheet, table in mapping.items():
            if sheet not in wb.sheetnames:
                continue
            rows = wb[sheet].iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            columns = [str(h) if h is not None else None for h in header]
            data = []
            for values in rows:
                if values is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                    continue
                data.append({c: v for c, v in zip(columns, values) if c is not None})
            tables[table] = data
        return tables
    finally:
        wb.close()


def apply_supplement(doc: XwpPlanDocument, supplement: Supplement, placeholder_ids=frozenset()) -> MergeResult:
    """Merge the supplement into ``doc`` (in place); see the module docstring for the rules.

    ``placeholder_ids``: ids of the template's placeholders (``build_plan`` passes them).
    """
    result = MergeResult(issues=list(supplement.issues))
    index = doc.by_id()
    plans = doc.of_type("WP_Plan")
    areas = doc.of_type("WP_Bereich")
    for obj in supplement.objects:
        name = type(obj).__xwp__.name
        where = f"Ergänzung {name}[{obj.id}]"
        target = index.get(obj.id)
        if target is not None:
            if type(target) is not type(obj):
                result.issues.append(Issue(where, "type", f"{where}: uuid gehört zu {type(target).__name__}, nicht übernommen."))
                result.skipped += 1
                continue
            changed = False
            for attr in type(obj).__xwp__.attrs:
                value = getattr(obj, attr.name)
                if value is None or value == []:
                    continue
                if attr.kind == REF and attr.many and attr.reverse:
                    continue  # reverse lists are rebuilt from the other side
                if not _same(getattr(target, attr.name), value, attr):
                    setattr(target, attr.name, value)
                    changed = True
            if changed:
                result.overridden += 1
                result.overridden_ids.add(obj.id)
            continue
        if name == "WP_Plan":
            result.issues.append(Issue(where, "type", f"{where}: ein zweiter WP_Plan wird nicht übernommen."))
            result.skipped += 1
            continue
        info = type(obj).__xwp__
        for attr in info.attrs:
            if attr.kind != REF or getattr(obj, attr.name) is not None:
                continue
            if attr.type == "WP_Plan" and plans:
                setattr(obj, attr.name, Ref.to(plans[0]))
            elif attr.type == "WP_Bereich" and areas:
                setattr(obj, attr.name, Ref.to(areas[0]))
        missing = [i for i in validate(obj, where) if i.code == "missing"]
        if missing and obj.id in placeholder_ids:
            result.open_placeholders += 1
            continue
        if missing:
            names = ", ".join(sorted({i.path.rsplit(".", 1)[-1] for i in missing}))
            result.issues.append(Issue(where, "missing", f"{where}: unvollständig ({names} fehlt), nicht übernommen."))
            result.skipped += 1
            continue
        doc.add(obj)
        index[obj.id] = obj
        result.added += 1
        result.added_ids.add(obj.id)
    doc.link_reverse()
    return result


def _same(current, new, attr) -> bool:
    """Whether a value read back equals the plan's value (references by id, geometry by shape)."""
    from shapely.geometry.base import BaseGeometry

    if attr.kind == REF:
        ids = lambda v: [r.id for r in v] if isinstance(v, list) else (v.id if v is not None else None)  # noqa: E731
        return ids(current) == ids(new)
    if isinstance(current, BaseGeometry) and isinstance(new, BaseGeometry):
        return current.equals_exact(new, 1e-3) or current.equals(new)
    try:
        return current == new
    except ValueError:  # e.g. comparing arrays
        return False
