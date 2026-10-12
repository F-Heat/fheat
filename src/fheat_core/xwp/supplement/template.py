"""Supplement template: what F|Heat cannot know, filled in by the planner.

Two files, both generated from the model metadata:

* ``xwp_ergaenzung.gpkg``: every table of the F|Heat layout (encode/gpkg.py),
  pre-filled with what the planner should check or complete: the plan, the
  proposed supply areas and suitability checks (without the computed values,
  so a new calculation is not overridden by an old file) and placeholders
  for the heat sources of the networks and for large consumers.
* ``xwp_ergaenzung.xlsx``: the tables without geometry (measures, texts,
  energy balances, aggregated figures and their child tables) for planners
  who prefer a spreadsheet, with code list sheets and drop-downs.

The reader (supplement/reader.py) accepts both, or either.
"""
from __future__ import annotations

import copy
from pathlib import Path
from typing import Iterable

from fheat_core.xwp.document import XwpPlanDocument
from fheat_core.xwp.encode.gpkg import TableDef, gpkg_schema, table_base, to_tables, write_gpkg
from fheat_core.xwp.model.base import ENUM, XwpFeature, model_module

# feature types the planner reviews; attributes computed by F|Heat are left empty in the template
REVIEW = {
    "WP_Plan": {"raeumlicherGeltungsbereich"},
    "WP_WaermeversorgungsgebietZieljahr": {"waermeMenge"},
    "WP_Eignungspruefung": set(),
}
INDEX_SHEET = "_tabellen"
NOTES_SHEET = "Hinweise"
MAX_SHEET_NAME = 31

INSTRUCTIONS = [
    "XWärmeplan-Ergänzung für F|Heat",
    "",
    "Hier tragen Sie ein, was F|Heat nicht aus seinen Daten ableiten kann.",
    "Jedes Tabellenblatt entspricht einer Tabelle der GeoPackage-Vorlage (Zuordnung im Blatt _tabellen).",
    "Zeile 1 enthält die Spaltennamen; nicht ändern. Erläuterungen stehen als Kommentar an den Spaltenköpfen.",
    "uuid leer lassen: F|Heat vergibt sie. Eine vorhandene uuid überschreibt das gleichnamige Objekt von F|Heat.",
    "gehoertzuplan_fk / gehoertzubereich_fk leer lassen: F|Heat setzt Plan und Bereich ein.",
    "Kindtabellen (z. B. wp_energiebilanz_eegesamtnachenergietraeger) verweisen über <tabelle>_id auf die uuid",
    "bzw. _id der übergeordneten Zeile. Codes (Spalten *_fkcl) stehen in den Codelisten-Blättern wp_*.",
    "Messwerte in der Einheit des Spaltenkommentars (z. B. MWh/a), Datum als JJJJ-MM-TT.",
]


def template_document(doc: XwpPlanDocument, placeholders: Iterable[XwpFeature] = ()) -> XwpPlanDocument:
    """The objects of ``doc`` the planner reviews (computed values removed), plus ``placeholders``."""
    out = XwpPlanDocument(id=doc.id)
    for f in doc.features:
        name = type(f).__xwp__.name
        if name not in REVIEW:
            continue
        c = copy.deepcopy(f)
        for attr in REVIEW[name]:
            setattr(c, attr, None)
        out.add(c)
    out.add(*placeholders)
    return out


def write_template(
    doc: XwpPlanDocument,
    gpkg_path: str | Path,
    xlsx_path: str | Path | None = None,
    placeholders: Iterable[XwpFeature] = (),
) -> dict[str, Path]:
    """Write the supplement template(s); returns the written paths by format."""
    template = template_document(doc, placeholders)
    written = {"gpkg": write_gpkg(template, gpkg_path, include_empty=True, check=False)}
    if xlsx_path is not None:
        written["xlsx"] = write_template_xlsx(template, xlsx_path)
    return written


def spreadsheet_tables(model) -> list[TableDef]:
    """Tables of the Excel template: feature tables without geometry and their child tables."""
    tables = gpkg_schema(model)
    aspatial = {table_base(t.feature_type) for t in tables if t.feature_type and t.geometry_type is None}

    def root(t: TableDef) -> str:
        return t.name if t.feature_type else _root_of(t, {x.name: x for x in tables})

    return [t for t in tables if (t.feature_type or t.parent) and root(t) in aspatial]


def _root_of(t: TableDef, by_name: dict[str, TableDef]) -> str:
    while t.parent is not None and t.parent in by_name and by_name[t.parent].parent is not None:
        t = by_name[t.parent]
    return t.parent or t.name


def sheet_names(names: list[str]) -> dict[str, str]:
    """Excel sheet name (max. 31 characters, unique) per table name."""
    out: dict[str, str] = {}
    used: set[str] = set()
    for name in names:
        sheet = name[:MAX_SHEET_NAME]
        n = 1
        while sheet in used:
            suffix = f"~{n}"
            sheet = name[: MAX_SHEET_NAME - len(suffix)] + suffix
            n += 1
        used.add(sheet)
        out[name] = sheet
    return out


def write_template_xlsx(template: XwpPlanDocument, path: str | Path, model=None) -> Path:
    """The tables without geometry as workbook (needs openpyxl, extra ``[excel]``)."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)
    if model is None:
        if template.features:
            model = model_module(template.features[0])
        else:
            from fheat_core.xwp.model import current as model
    rows_by_table = {t.name: t.rows for t in to_tables(template, include_empty=True, check=False)} if template.features else {}
    tables = spreadsheet_tables(model)
    code_lists = {table_base(name): enum_cls for name, enum_cls in model.ENUMS.items()}
    names = sheet_names([t.name for t in tables] + sorted(code_lists))
    wb = Workbook()
    notes = wb.active
    notes.title = NOTES_SHEET
    for i, line in enumerate(INSTRUCTIONS, start=1):
        notes.cell(row=i, column=1, value=line)
    notes["A1"].font = Font(bold=True, size=14)
    index = wb.create_sheet(INDEX_SHEET)
    index.append(["blatt", "tabelle", "inhalt"])
    bold = Font(bold=True)

    for name, enum_cls in sorted(code_lists.items()):
        ws = wb.create_sheet(names[name])
        ws.append(["code", "name", "label"])
        for e in enum_cls:
            ws.append([e.code, e.name, e.label])
        index.append([names[name], name, "Codeliste (nur zum Nachschlagen)"])
        for cell in ws[1]:
            cell.font = bold

    for t in tables:
        ws = wb.create_sheet(names[t.name])
        ws.append([c.name for c in t.columns])
        for col, c in enumerate(t.columns, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font = bold
            cell.comment = Comment(f"{c.alias}\n{c.doc}"[:2000], "F|Heat")
        for row in rows_by_table.get(t.name, []):
            ws.append([row.get(c.name) for c in t.columns])
        attrs = {a.name.lower() + "_fkcl": a for a in _enum_attrs(t, model)}
        for col, c in enumerate(t.columns, start=1):
            attr = attrs.get(c.name)
            if attr is None:
                continue
            sheet = names[table_base(attr.type)]
            n = len(model.ENUMS[attr.type])
            dv = DataValidation(type="list", formula1=f"='{sheet}'!$A$2:$A${n + 1}", allow_blank=True)
            dv.error = f"Bitte einen Code aus {attr.type} wählen."
            ws.add_data_validation(dv)
            letter = ws.cell(row=1, column=col).column_letter
            dv.add(f"{letter}2:{letter}1000")
        index.append([names[t.name], t.name, t.description[:200] if t.description else t.identifier])
    for cell in index[1]:
        cell.font = bold
    wb.save(path)
    return path


def _enum_attrs(t: TableDef, model):
    if t.feature_type:
        return [a for a in model.CLASSES[t.feature_type].__xwp__.attrs if a.kind == ENUM and not a.many]
    attr = t.attr
    if attr.kind == ENUM:
        return [attr]
    if attr.kind == "datatype":
        return [a for a in model.CLASSES[attr.type].__xwp__.attrs if a.kind == ENUM and not a.many]
    return []


__all__ = ["REVIEW", "spreadsheet_tables", "template_document", "write_template", "write_template_xlsx"]
