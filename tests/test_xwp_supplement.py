"""Supplement file and readers (phase 8): GML and GeoPackage back into the model, template, merge."""
from __future__ import annotations

import datetime
from pathlib import Path

import geopandas as gpd
import pytest
from shapely.geometry import Point

from fheat_core.config import XwpConfig, XwpGemeinde, XwpPlanConfig
from fheat_core.state import PipelineState
from fheat_core.xwp.decode.tables import from_tables
from fheat_core.xwp.document import XwpPlanDocument
from fheat_core.xwp.encode.jsonfg import to_jsonfg
from fheat_core.xwp.mapping import NetworkInput, build_plan
from fheat_core.xwp.model import Measure, validate
from fheat_core.xwp.model import current as m
from fheat_core.xwp.supplement import apply_supplement, read_supplement, write_template
from tests.xwp_synth import X0, Y0, document_all, small_town

pytest.importorskip("lxml")
pytest.importorskip("pyogrio")

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "xwp"
CRS = "EPSG:25832"

# classes F|Heat does not fill automatically (Modellreferenz "E" plus those whose mandatory
# attributes F|Heat cannot know, see docs/xwaermeplan/abbildung_planung.md)
SUPPLEMENT_CLASSES = [
    "WP_Massnahme", "WP_TextAbschnitt", "WP_WaermeversorgungsgebietUmsetzung", "WP_Energieeinsparpotenzialgebiet",
    "WP_NichtBeplantesTeilgebiet", "WP_PotenzialWaermespeicherung", "WP_Ausschlussgebiet", "WP_Anschlusszwang",
    "WP_AnschlussGruenesMethan", "WP_WaermeKaelteNetz", "WP_Kaeltenetz", "WP_Kaeltelinie", "WP_Kaelteversorgungsgebiet",
    "WP_Energiebilanz", "WP_PotenzialanalyseAggregiert", "WP_Grossverbraucher", "WP_PotenzialWaermenutzung",
    "WP_EnergieeinsparungGebaeude", "WP_WaermeversorgungsgebietStuetzjahr",
]


def comparable(doc):
    """Ids, types, EPSG:25832 geometry and properties: everything a plan states."""
    data = to_jsonfg(doc, check=False)
    return {f["id"]: (f["featureType"], f["place"], f["properties"]) for f in data["features"]}


# --------------------------------------------------------------------------
# GML reader


@pytest.mark.parametrize("maximal", [False, True], ids=["minimal", "maximal"])
@pytest.mark.parametrize("variable_as", ["surface", "point"])
def test_gml_round_trip_of_every_class(maximal, variable_as):
    from fheat_core.xwp.decode.gml import read_gml
    from fheat_core.xwp.encode.gml import to_gml

    doc, _ = document_all(m, maximal, variable_as)
    back, issues = read_gml(to_gml(doc))
    back.link_reverse()
    assert issues == [] and back.id == doc.id
    assert comparable(back) == comparable(doc)


@pytest.mark.parametrize("name", sorted(p.name for p in FIXTURES.glob("*.gml")))
def test_reference_plans_read_write_read(name):
    from fheat_core.xwp.decode.gml import read_gml
    from fheat_core.xwp.encode.gml import to_gml, validate_gml

    first, issues = read_gml(FIXTURES / name)
    assert issues == []
    assert all(i.code == "uom" for i in first.errors())  # only the unit deviation G1 of abgleich.md
    tree = to_gml(first, check=False)
    assert validate_gml(tree) == []
    second, _ = read_gml(tree)
    assert comparable(second) == comparable(first)


def test_xmas_and_ldproxy_encodings_mean_the_same():
    from fheat_core.xwp.decode.gml import read_gml

    xmas = comparable(read_gml(FIXTURES / "rheinberg_xmas_auszug.gml")[0])
    ldproxy = comparable(read_gml(FIXTURES / "rheinberg_ldproxy_auszug.gml")[0])
    common = set(xmas) & set(ldproxy)
    assert len(common) > 30
    assert all(xmas[k] == ldproxy[k] for k in common)


def test_foreign_units_are_kept_and_reported():
    from fheat_core.xwp.decode.gml import read_gml

    doc, _ = read_gml(FIXTURES / "rheinberg_xmas_auszug.gml")
    bilanz = next(f for f in doc.of_type("WP_Energiebilanz") if f.eeGesamtGebaeude is not None)
    assert bilanz.eeGesamtGebaeude == Measure(bilanz.eeGesamtGebaeude.value, "MWh/a")
    assert [i.code for i in validate(bilanz) if i.path.endswith("eeGesamtGebaeude")] == ["uom"]


def test_other_crs_and_unknown_elements():
    from fheat_core.xwp.decode.gml import read_gml

    gml = f"""<sf:FeatureCollection xmlns="{m.NAMESPACE}" xmlns:gml="http://www.opengis.net/gml/3.2"
        xmlns:xlink="http://www.w3.org/1999/xlink" xmlns:sf="http://www.opengis.net/ogcapi-features-1/1.0/sf">
      <sf:featureMember><WP_Grossverbraucher gml:id="g1">
        <gehoertZuBereich xlink:href="#b1"/>
        <position><gml:Point srsName="urn:ogc:def:crs:EPSG::4326"><gml:pos>52.0 7.5</gml:pos></gml:Point></position>
        <status>1000</status><energiebedarf>1000</energiebedarf><energieeinsatz>1000</energieeinsatz>
        <versorgungsmedium>1000</versorgungsmedium><gibtEsNicht>1</gibtEsNicht>
      </WP_Grossverbraucher></sf:featureMember>
    </sf:FeatureCollection>""".encode()
    doc, issues = read_gml(gml)
    [big] = doc.features
    assert 395_000 < big.position.x < 405_000 and 5_755_000 < big.position.y < 5_765_000  # lat/lon order respected
    assert [i.code for i in issues] == ["attribute", "ref"]  # unknown element, dangling link


# --------------------------------------------------------------------------
# GeoPackage reader and table rows


@pytest.mark.parametrize("maximal", [False, True], ids=["minimal", "maximal"])
def test_gpkg_round_trip_of_every_class(tmp_path, maximal):
    from fheat_core.xwp.decode.gpkg import read_gpkg
    from fheat_core.xwp.encode.gpkg import write_gpkg

    doc, _ = document_all(m, maximal)
    features, issues = read_gpkg(write_gpkg(doc, tmp_path / "plan.gpkg"))
    back = XwpPlanDocument(features=features, id=doc.id)
    back.link_reverse()
    assert issues == []
    assert comparable(back) == comparable(doc)


def test_table_rows_report_problems_by_row():
    tables = {
        "wp_massnahme": [
            {"uuid": None, "bezeichnung": "Quartierskonzept", "kategorie_fkcl": "9", "abschlussumsetzung": "2030-13-01"},
        ],
        "wp_massnahme_kostentragendersektor": [{"_id": None, "wp_massnahme_id": "nicht-da", "kostentragendersektor_fkcl": "1000"}],
        "wp_gibtsnicht": [{"x": 1}],
    }
    tables["wp_massnahme"][0]["spalte_falsch"] = 1
    objects, issues = from_tables(tables, m)
    [massnahme] = objects
    assert massnahme.bezeichnung == "Quartierskonzept" and massnahme.kategorie is None
    messages = " | ".join(i.message for i in issues)
    assert "wp_massnahme Zeile 1: kategorie erwartet einen Code aus WP_KategorieMassnahme" in messages
    assert "abschlussUmsetzung erwartet ein Datum" in messages
    assert "wp_massnahme_kostentragendersektor Zeile 1: Zeile verweist auf keine Zeile" in messages
    assert "Tabelle wp_gibtsnicht" in messages and "Unbekannte Spalten in wp_massnahme: spalte_falsch" in messages


def test_table_values_are_coerced():
    tables = {"wp_massnahme": [{"uuid": "6F1C1B8E-0000-4000-8000-000000000001", "bezeichnung": "A",
                                "kategorie_fkcl": 1000.0, "abschlussumsetzung": datetime.datetime(2030, 1, 1, 0, 0),
                                "kosten": "12500", "gehoertzuplan_fk": "6f1c1b8e-0000-4000-8000-000000000002"}]}
    [obj], issues = from_tables(tables, m)
    assert issues == []
    assert obj.id == "6f1c1b8e-0000-4000-8000-000000000001"
    assert obj.kategorie.code == "1000" and obj.abschlussUmsetzung == datetime.date(2030, 1, 1)
    assert obj.gehoertZuPlan.id.endswith("0002") and obj.kosten == 12500


# --------------------------------------------------------------------------
# template, reading and merging


def _plan(supplement=None, **cfg):
    buildings, wld, streets, boundary = small_town()
    plan_cfg = XwpPlanConfig(name="Wärmeplan Test", gemeinde=XwpGemeinde("Teststadt", "051234567890"),
                             beschluss_datum=datetime.date(2026, 6, 30))
    source = gpd.GeoDataFrame(geometry=[Point(X0 + 5, Y0 + 5)], crs=CRS)
    state = PipelineState(buildings_gdf=buildings, net_gdf=None, source_gdf=source)
    return build_plan(buildings, wld, XwpConfig(plan=plan_cfg, privacy_min_buildings=1, **cfg),
                      networks=[NetworkInput(state, name="Netz West")], streets=streets, boundary=boundary,
                      data_year=2024, supplement=supplement)


def test_template_holds_review_objects_and_placeholders(tmp_path):
    import pyogrio

    doc, notes = _plan(large_consumer_mwh=50)
    assert [type(p).__name__ for p in notes.placeholders] == ["WP_PotenzialWaermenutzung", "WP_Grossverbraucher"]
    paths = write_template(doc, tmp_path / "xwp_ergaenzung.gpkg", tmp_path / "xwp_ergaenzung.xlsx", notes.placeholders)
    layers = dict(pyogrio.list_layers(paths["gpkg"]))
    assert len(layers) == 92  # every table, for the planner to fill in
    areas = pyogrio.read_dataframe(paths["gpkg"], layer="wp_waermeversorgungsgebietzieljahr")
    assert len(areas) == 2 and areas["waermemenge"].isna().all()  # computed values left out
    assert len(pyogrio.read_dataframe(paths["gpkg"], layer="wp_waermelinie")) == 0  # F|Heat's own results not offered
    sources = pyogrio.read_dataframe(paths["gpkg"], layer="wp_potenzialwaermenutzung_punkt")
    assert len(sources) == 1 and "Energieträger" in sources.loc[0, "hinweis"]

    from openpyxl import load_workbook

    wb = load_workbook(paths["xlsx"])
    assert {"Hinweise", "_tabellen", "wp_massnahme", "wp_energiebilanz", "wp_status"} <= set(wb.sheetnames)
    assert all(len(name) <= 31 for name in wb.sheetnames)
    sheet = wb["wp_massnahme"]
    assert sheet["A1"].value == "uuid" and sheet["A1"].comment is not None
    assert any("WP_KategorieMassnahme" in (dv.error or "") for dv in sheet.data_validations.dataValidation)


def _edit(path, layer, geometry_type, change):
    """Change one table of a GeoPackage the way a planner would in QGIS."""
    import pyogrio

    frame = change(pyogrio.read_dataframe(path, layer=layer))
    pyogrio.write_dataframe(frame, path, layer=layer, geometry_type=geometry_type, layer_options={"OVERWRITE": "YES"})


def test_supplement_overrides_adds_and_skips(tmp_path):
    doc, notes = _plan()
    paths = write_template(doc, tmp_path / "e.gpkg", tmp_path / "e.xlsx", notes.placeholders)
    area_id = doc.of_type("WP_WaermeversorgungsgebietZieljahr")[0].id
    # the planner decides on the supply type of one proposal and names the source's energy carrier
    def decide(frame):
        frame.loc[frame["uuid"] == area_id, "versorgungsart_fkcl"] = "10003"
        return frame

    _edit(paths["gpkg"], "wp_waermeversorgungsgebietzieljahr", "MultiPolygon", decide)
    _edit(paths["gpkg"], "wp_potenzialwaermenutzung_punkt", "MultiPoint",
          lambda f: f.assign(energietraeger_fkcl="8000", waermepotenzial=4200.0))

    from openpyxl import load_workbook

    wb = load_workbook(paths["xlsx"])
    sheet = wb["wp_massnahme"]
    header = [c.value for c in sheet[1]]
    row = {"kategorie_fkcl": "2000", "bezeichnung": "Wärmenetz West", "abschlussumsetzung": "2030-12-31"}
    sheet.append([row.get(h) for h in header])
    wb.save(paths["xlsx"])
    supplement = read_supplement(gpkg=paths["gpkg"], xlsx=paths["xlsx"])
    assert supplement.issues == []

    merged, notes = _plan(supplement=supplement)
    assert merged.errors() == []
    area = next(a for a in merged.of_type("WP_WaermeversorgungsgebietZieljahr") if a.id == area_id)
    assert area.versorgungsart.code == "10003" and area.waermeMenge is not None  # override, computed value kept
    [source] = merged.of_type("WP_PotenzialWaermenutzung")
    assert source.energietraeger.code == "8000" and source.waermepotenzial == 4200.0
    [measure] = merged.of_type("WP_Massnahme")
    plan = merged.of_type("WP_Plan")[0]
    assert measure.gehoertZuPlan.id == plan.id and [r.id for r in plan.massnahme] == [measure.id]
    assert notes.supplement_added == 2 and notes.supplement_overridden >= 1
    assert any(msg.startswith("Ergänzung:") for msg in notes.messages)


def test_incomplete_placeholder_is_not_taken_over(tmp_path):
    doc, notes = _plan()
    paths = write_template(doc, tmp_path / "e.gpkg", None, notes.placeholders)
    merged, notes = _plan(supplement=read_supplement(gpkg=paths["gpkg"]))
    assert merged.of_type("WP_PotenzialWaermenutzung") == []
    assert notes.supplement_skipped == 1
    assert any("energietraeger fehlt" in i.message for i in notes.supplement_issues)


def test_a_second_plan_is_refused():
    doc, _ = _plan()
    other = m.WP_Plan(name="Anderer Plan")
    result = apply_supplement(doc, _supplement([other]))
    assert result.skipped == 1 and len(doc.of_type("WP_Plan")) == 1


def _supplement(objects):
    from fheat_core.xwp.supplement import Supplement

    return Supplement(objects=objects)


def test_every_supplement_class_can_be_filled_in(tmp_path):
    """Each class F|Heat does not fill can come complete from the supplement files."""
    from fheat_core.xwp.encode.gpkg import write_gpkg
    from fheat_core.xwp.supplement.template import spreadsheet_tables, write_template_xlsx

    doc, objects = document_all(m, True)
    wanted = {name: objects[name] for name in SUPPLEMENT_CLASSES}
    target = XwpPlanDocument(features=[f for f in doc.features if f not in wanted.values()], id=doc.id)
    for f in target.features:  # drop the plan's references to what the supplement brings
        for attr in type(f).__xwp__.attrs:
            if attr.kind == "ref" and attr.many and attr.reverse:
                setattr(f, attr.name, [r for r in getattr(f, attr.name) if r.id not in {o.id for o in wanted.values()}])
    spreadsheet = {t.feature_type for t in spreadsheet_tables(m) if t.feature_type}
    gpkg = write_gpkg(XwpPlanDocument(features=[o for n, o in wanted.items() if n not in spreadsheet]),
                      tmp_path / "e.gpkg", check=False)
    xlsx = write_template_xlsx(XwpPlanDocument(features=[o for n, o in wanted.items() if n in spreadsheet]),
                               tmp_path / "e.xlsx", model=m)
    supplement = read_supplement(gpkg=gpkg, xlsx=xlsx)
    assert supplement.issues == []
    result = apply_supplement(target, supplement)
    assert result.issues == [] and result.added == len(SUPPLEMENT_CLASSES)
    assert target.errors() == []
    assert comparable(target) == comparable(doc)
